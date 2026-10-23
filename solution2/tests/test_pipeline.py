"""End-to-end pipeline tests: a scripted extractor stands in for the VLM, the text layer for the crop reader."""
import hashlib
import json
from pathlib import Path

import openpyxl
import pytest

from solution2.harness.corpus import ROOT, load_corpus
from solution2.pipeline import run_pipeline
from solution2.repair import TextLayerReader

PDF = ROOT / "balance_sheet_after_correction.pdf"
MISSED = ROOT / "benchmark_results/qwen3-vl_4b-instruct/pass1_raw_output.json"      # 建設仮勘定 4,553 read as '－'
YAMATO_PDF = ROOT / "dataset/benchmark_suite/company_01_yamato/balance_sheet.pdf"
YAMATO_BAD = ROOT / "solution2/fixtures/real_vlm/yamato_p1__qwen3-vl_4b-instruct.json"
pytestmark = pytest.mark.skipif(not (PDF.exists() and MISSED.exists()), reason="fixtures not available")

text_reader = lambda page: TextLayerReader(page)


def snapshot(folder: Path):
    return sorted((str(p.relative_to(folder)), p.stat().st_size, p.stat().st_mtime_ns)
                  for p in folder.rglob("*") if p.is_file() and "__pycache__" not in p.parts)


def run(tmp_path, extractor, pdf=PDF, **kw):
    return run_pipeline(pdf, tmp_path, extractor=extractor, reader_factory=text_reader, **kw)


def report(tmp_path, page=1):
    return json.loads((tmp_path / f"page{page}.json").read_text(encoding="utf-8"))


def test_missed_value_is_repaired_and_every_cell_carries_its_status(tmp_path):
    out = run(tmp_path, lambda img: MISSED.read_text(encoding="utf-8"))[0]
    assert out.status == "repaired" and out.trusted and out.vlm_calls == 3
    assert {"json", "xlsx", "md", "vlm_raw", "stage1"} <= set(out.files)
    rep = report(tmp_path)
    assert rep["trusted"] and rep["unit_multiplier"] == 1000 and rep["coverage"] == 1.0
    cell = next(r for r in rep["rows"] if r["label"] == "建設仮勘定")["cells"][1]
    assert (cell["raw"], cell["value"], cell["value_jpy"]) == ("4,553", 4553, 4_553_000)
    assert cell["repaired"] and cell["original_raw"] == "－" and cell["status"] in ("exact", "rounding")
    assert sum(c["repaired"] for r in rep["rows"] for c in r["cells"]) == 1               # nothing else touched
    assert all(c["status"] in ("exact", "rounding") for r in rep["rows"] for c in r["cells"])
    assert rep["repairs"][0]["applied"] and rep["periods"][1]["date"] == "2026-03-31"


def test_what_the_vlm_returned_is_kept_byte_for_byte(tmp_path):
    original = MISSED.read_text(encoding="utf-8")
    run(tmp_path, lambda img: original)
    assert (tmp_path / "page1_vlm_raw.json").read_text(encoding="utf-8") == original


def test_excel_marks_the_repaired_cell_and_has_a_verification_sheet(tmp_path):
    run(tmp_path, lambda img: MISSED.read_text(encoding="utf-8"))
    wb = openpyxl.load_workbook(tmp_path / "page1.xlsx")
    ws, vs = wb["Balance Sheet"], wb["Verification"]
    assert "REPAIRED" in ws["A2"].value
    commented = [c for row in ws.iter_rows() for c in row if c.comment]
    assert len(commented) == 1 and "'－'" in commented[0].comment.text
    assert any("applied" in str(cell.value) for row in vs.iter_rows() for cell in row)


def test_clean_page_needs_no_repair(tmp_path):
    gold = load_corpus()["baseline"]
    out = run(tmp_path, lambda img: gold.model_dump_json())[0]
    assert out.status == "trusted" and out.vlm_calls == 0
    assert not any(c["repaired"] for r in report(tmp_path)["rows"] for c in r["cells"])


def test_verify_only_mode_reports_the_failure_without_repairing(tmp_path):
    out = run(tmp_path, lambda img: MISSED.read_text(encoding="utf-8"), do_repair=False)[0]
    assert out.status == "needs_review" and not out.trusted and out.vlm_calls == 0
    rep = report(tmp_path)
    assert any(f["kind"] == "equation_failed" for f in rep["findings"]) and rep["needs_review"]
    assert any(c["status"] == "failed" for r in rep["rows"] for c in r["cells"])
    assert "NOT TRUSTED" in openpyxl.load_workbook(tmp_path / "page1.xlsx")["Balance Sheet"]["A2"].value


@pytest.mark.skipif(not (YAMATO_PDF.exists() and YAMATO_BAD.exists()), reason="Yamato fixture not available")
def test_a_two_column_page_is_exported_for_review_not_guessed(tmp_path):
    out = run(tmp_path, lambda img: YAMATO_BAD.read_text(encoding="utf-8"), pdf=YAMATO_PDF)[0]
    assert out.status == "needs_review" and out.vlm_calls == 0
    rep = report(tmp_path)
    assert not rep["trusted"] and rep["repairs"][0]["outcome"] == "no_region" and "two-column" in rep["repairs"][0]["detail"]


def test_unusable_vlm_output_is_a_reported_failure(tmp_path):
    assert run(tmp_path, lambda img: None)[0].status == "extraction_failed"
    assert run(tmp_path / "b", lambda img: "this is not json")[0].status == "extraction_failed"
    assert report(tmp_path / "b")["trusted"] is False


def test_a_page_without_a_table_is_skipped_not_failed(tmp_path):
    blank = json.dumps({"metadata": {"title": "t", "unit_raw": "", "periods": []},
                        "rows": [{"row_index": 1, "indent_level": 0, "account_raw": "表紙", "role": "section_header"}]})
    out = run(tmp_path, lambda img: blank)[0]
    assert out.status == "no_table" and not out.trusted


def test_solution1_is_left_untouched(tmp_path):
    folder = ROOT / "solution1"
    before = snapshot(folder)
    run(tmp_path, lambda img: MISSED.read_text(encoding="utf-8"))
    assert snapshot(folder) == before


# ----------------------------------------------------------------------------- statements that continue across pages
SILVER_NIKON_P2 = ROOT / "solution2/fixtures/silver/silver_nikon_p2.json"
CALBEE_PDF = ROOT / "dataset/benchmark_suite/company_02_calbee/balance_sheet.pdf"


def halves(corrupt=False):
    """One clean statement cut mid-section, as a page break does: neither half balances alone."""
    from solution1.models import RawDocumentTree
    whole = RawDocumentTree.model_validate_json(SILVER_NIKON_P2.read_text(encoding="utf-8"))
    first = whole.model_copy(update={"rows": whole.rows[:12]})
    second_rows = [r.model_copy(update={"row_index": k + 1}) for k, r in enumerate(whole.rows[12:])]
    if corrupt:
        second_rows[2] = second_rows[2].model_copy(update={"values": ["999,999", *second_rows[2].values[1:]]})
    return first.model_dump_json(), whole.model_copy(update={"rows": second_rows}).model_dump_json()


def two_page_extractor(texts):
    calls = iter(texts)
    return lambda img: next(calls)


@pytest.mark.skipif(not (SILVER_NIKON_P2.exists() and CALBEE_PDF.exists()), reason="fixtures not available")
def test_a_split_statement_is_verified_jointly_with_no_repair_calls(tmp_path):
    outs = run(tmp_path, two_page_extractor(halves()), pdf=CALBEE_PDF)
    assert [o.status for o in outs] == ["trusted_joined", "trusted_joined"] and all(o.trusted and o.vlm_calls == 0 for o in outs)
    assert outs[0].joined_with == [2] and outs[1].joined_with == [1]
    rep = report(tmp_path, 2)
    assert rep["trusted"] and rep["status"] == "trusted_joined" and rep["source"]["verified_jointly_with_pages"] == [1]
    assert rep["coverage"] is not None and all(c["status"] in ("exact", "rounding") for r in rep["rows"] for c in r["cells"])


@pytest.mark.skipif(not (SILVER_NIKON_P2.exists() and CALBEE_PDF.exists()), reason="fixtures not available")
def test_without_joining_each_half_is_flagged_as_before(tmp_path):
    outs = run(tmp_path, two_page_extractor(halves()), pdf=CALBEE_PDF, join=False, do_repair=False)
    assert [o.status for o in outs] == ["needs_review", "needs_review"] and not any(o.trusted for o in outs)


@pytest.mark.skipif(not (SILVER_NIKON_P2.exists() and CALBEE_PDF.exists()), reason="fixtures not available")
def test_joining_never_launders_a_real_error_into_trust(tmp_path):
    outs = run(tmp_path, two_page_extractor(halves(corrupt=True)), pdf=CALBEE_PDF, do_repair=False)
    assert [o.status for o in outs] == ["needs_review", "needs_review"] and not any(o.joined_with for o in outs)


@pytest.mark.skipif(not (SILVER_NIKON_P2.exists() and CALBEE_PDF.exists()), reason="fixtures not available")
def test_pages_with_different_columns_are_not_joined(tmp_path):
    first, second = halves()
    second = second.replace("（単位：百万円）", "（単位：千円）").replace("(単位：百万円)", "(単位：千円)")
    outs = run(tmp_path, two_page_extractor([first, second]), pdf=CALBEE_PDF, do_repair=False)
    assert not any(o.joined_with for o in outs)


@pytest.mark.skipif(not (SILVER_NIKON_P2.exists() and CALBEE_PDF.exists()), reason="fixtures not available")
def test_every_stage_is_timed_and_the_total_adds_up(tmp_path):
    outs = run(tmp_path, two_page_extractor(halves()), pdf=CALBEE_PDF)
    for o in outs:
        t = o.timing
        assert {"render_s", "extract_s", "verify_s", "join_s", "export_s", "total_s"} <= set(t)       # both halves were flagged alone
        assert t["total_s"] == pytest.approx(sum(v for k, v in t.items() if k != "total_s"), rel=1e-6)
        assert all(v >= 0 for v in t.values())
    rep = report(tmp_path, 1)
    assert rep["timing"]["total_s"] > 0 and "extract_s" in rep["timing"]
    summary = json.loads((tmp_path / "summary.json").read_text(encoding="utf-8"))
    assert summary[0]["timing"]["total_s"] > 0 and summary[0]["joined_with"] == [2]


@pytest.mark.skipif(not (SILVER_NIKON_P2.exists() and CALBEE_PDF.exists()), reason="fixtures not available")
def test_a_repaired_half_is_rejoined_with_its_neighbour(tmp_path, monkeypatch):
    """The first joint check fails because a half has a wrong value; repair fixes it, but a half can never be trusted alone,
    so the pages must be joined again after repair."""
    from types import SimpleNamespace
    from solution2 import pipeline
    from solution2.verify import parse_stage1, verify_document
    clean_second = parse_stage1(halves()[1])[0]

    def fake_repair(raw, page, reader, config=None):          # puts back the value the VLM got wrong; a half still cannot balance alone
        tree = clean_second if any("999,999" in v for r in raw.rows for v in r.values) else raw
        return SimpleNamespace(raw=tree, verdict=verify_document(tree), attempts=[], vlm_calls=3, status="needs_review")

    monkeypatch.setattr(pipeline, "repair", fake_repair)
    outs = run(tmp_path, two_page_extractor(halves(corrupt=True)), pdf=CALBEE_PDF)
    assert [o.status for o in outs] == ["trusted_joined", "trusted_joined"] and all(o.trusted for o in outs)
    assert outs[0].joined_with == [2] and outs[1].joined_with == [1]
